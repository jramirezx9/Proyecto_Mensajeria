import tkinter as tk
from tkinter import ttk, filedialog, messagebox, scrolledtext
import pandas as pd
import webbrowser
import urllib.parse
import time
import threading
import re
import os
import smtplib
from email.mime.text import MIMEText
from email.mime.multipart import MIMEMultipart
from email.mime.base import MIMEBase
from email import encoders
from datetime import datetime

# Optional imports for reading Word/PDF files
try:
    import docx
except ImportError:
    docx = None
try:
    import PyPDF2
except ImportError:
    PyPDF2 = None

# dksy ppty xoht sodm
# ==============================================================================
#  HELPER: Personalize messages with row data
# ==============================================================================
def format_message(template, row):
    """Replaces placeholders like {name}, {phone} with actual row data"""
    msg = template
    for col in row.index:
        placeholder = "{" + str(col) + "}"
        msg = msg.replace(placeholder, str(row[col]))
    return msg


# ==============================================================================
#  SECONDARY WINDOW: Channel Configuration Popup
# ==============================================================================
class ChannelPopup(tk.Toplevel):
    def __init__(self, parent, app, channel_name):
        super().__init__(parent)
        self.app = app
        self.channel_name = channel_name
        self.title(f"Configure {channel_name} Messaging")
        self.geometry("950x650")
        self.configure(bg="#2C2C2C")
        self.grab_set()  # Modal window
        
        self.attachment_path = None
        self.setup_ui()
        
    def setup_ui(self):
        # Title
        tk.Label(self, text=f"📨 {self.channel_name} Configuration", 
                 font=("Arial", 14, "bold"), fg="#FFF", bg="#2C2C2C").pack(pady=10)
        
        # Main Frame
        main_frame = tk.Frame(self, bg="#2C2C2C")
        main_frame.pack(fill="both", expand=True, padx=10, pady=5)
        
        # ---------------- LEFT SIDE: Recipients List ----------------
        left_frame = tk.LabelFrame(main_frame, text="Filtered Recipients (Phone & Email Visible)", bg="#2C2C2C", fg="#FFF", font=("Arial", 10, "bold"))
        left_frame.pack(side="left", fill="both", expand=True, padx=5)
        
        self.tree = ttk.Treeview(left_frame, columns=("Num", "Name", "Phone", "Email", "Group"), show="headings", height=18)
        for col in ("Num", "Name", "Phone", "Email", "Group"):
            self.tree.heading(col, text=col)
            self.tree.column(col, width=100 if col != "Name" else 150)
        
        tree_scroll = ttk.Scrollbar(left_frame, orient="vertical", command=self.tree.yview)
        self.tree.configure(yscroll=tree_scroll.set)
        self.tree.pack(side="left", fill="both", expand=True)
        tree_scroll.pack(side="right", fill="y")
        
        self.load_recipients()
        
        # ---------------- RIGHT SIDE: Message & Settings ----------------
        right_frame = tk.Frame(main_frame, bg="#2C2C2C")
        right_frame.pack(side="right", fill="both", expand=True, padx=5)
        
        # File Selector
        file_frame = tk.Frame(right_frame, bg="#2C2C2C")
        file_frame.pack(fill="x", pady=5)
        tk.Label(file_frame, text="Message/Attachment File:", fg="#FFF", bg="#2C2C2C").pack(side="left")
        self.file_label = tk.Label(file_frame, text="None selected", fg="#FFA", bg="#2C2C2C")
        self.file_label.pack(side="left", padx=5)
        tk.Button(file_frame, text="Browse...", command=self.browse_file, bg="#4CAF50", fg="#FFF").pack(side="right")
        
        # Message Content (Editable Text Area)
        tk.Label(right_frame, text="Message Content (Use {name}, {phone}, {email} to personalize):", fg="#FFF", bg="#2C2C2C").pack(anchor="w")
        self.text_widget = scrolledtext.ScrolledText(right_frame, width=50, height=10, bg="#1E1E1E", fg="#FFF", insertbackground="#FFF")
        self.text_widget.pack(fill="both", expand=True, pady=5)
        
        # Email Specific Settings (Only if channel is Email)
        if self.channel_name == "Email":
            email_frame = tk.LabelFrame(right_frame, text="Email SMTP Settings", bg="#2C2C2C", fg="#FFF")
            email_frame.pack(fill="x", pady=5)
            
            tk.Label(email_frame, text="Provider:", fg="#FFF", bg="#2C2C2C").grid(row=0, column=0, sticky="w", padx=5, pady=2)
            self.provider_var = tk.StringVar(value="Gmail")
            ttk.Combobox(email_frame, textvariable=self.provider_var, values=["Gmail", "Outlook"], state="readonly", width=15).grid(row=0, column=1, padx=5, pady=2)
            
            tk.Label(email_frame, text="Sender Email:", fg="#FFF", bg="#2C2C2C").grid(row=1, column=0, sticky="w", padx=5, pady=2)
            self.email_entry = tk.Entry(email_frame, width=25)
            self.email_entry.grid(row=1, column=1, padx=5, pady=2)
            
            tk.Label(email_frame, text="App Password:", fg="#FFF", bg="#2C2C2C").grid(row=2, column=0, sticky="w", padx=5, pady=2)
            self.pass_entry = tk.Entry(email_frame, width=25, show="*")
            self.pass_entry.grid(row=2, column=1, padx=5, pady=2)
            
            tk.Label(email_frame, text="Subject:", fg="#FFF", bg="#2C2C2C").grid(row=3, column=0, sticky="w", padx=5, pady=2)
            self.subject_entry = tk.Entry(email_frame, width=25)
            self.subject_entry.insert(0, "Message for {name}")
            self.subject_entry.grid(row=3, column=1, padx=5, pady=2)
            
            self.attach_var = tk.BooleanVar(value=True)
            tk.Checkbutton(email_frame, text="Attach file to email", variable=self.attach_var, bg="#2C2C2C", fg="#FFF", selectcolor="#555").grid(row=4, column=0, columnspan=2, sticky="w", padx=5)

        # ---------------- BOTTOM: Action Buttons ----------------
        btn_frame = tk.Frame(self, bg="#2C2C2C")
        btn_frame.pack(fill="x", pady=10)
        
        tk.Button(btn_frame, text="🚀 PROCEED & SEND", command=self.proceed, 
                  bg="#2196F3", fg="#FFF", font=("Arial", 11, "bold"), width=20, height=2).pack(side="left", expand=True, padx=20)
        tk.Button(btn_frame, text="❌ CANCEL", command=self.destroy, 
                  bg="#F44336", fg="#FFF", font=("Arial", 11, "bold"), width=20, height=2).pack(side="right", expand=True, padx=20)

    def load_recipients(self):
        df = self.app.get_filtered_df()
        for _, row in df.iterrows():
            self.tree.insert("", "end", values=(
                row.get('num', ''), row.get('name', ''), row.get('phone', ''), 
                row.get('email', ''), row.get('group', '')
            ))
            
    def browse_file(self):
        path = filedialog.askopenfilename(
            filetypes=[
                ("All supported", "*.txt *.docx *.pdf *.mp3 *.wav"),
                ("Text", "*.txt"), ("Word", "*.docx"), ("PDF", "*.pdf"), 
                ("Audio", "*.mp3 *.wav"), ("All files", "*.*")
            ], parent=self
        )
        if path:
            self.attachment_path = path
            self.file_label.config(text=os.path.basename(path), fg="#66BB6A")
            
            # Extract text to populate the editable message box
            text_content = self.extract_text(path)
            self.text_widget.delete("1.0", tk.END)
            self.text_widget.insert(tk.END, text_content)
            
    def extract_text(self, path):
        ext = os.path.splitext(path)[1].lower()
        if ext == ".txt":
            with open(path, 'r', encoding='utf-8') as f:
                return f.read()
        elif ext == ".docx" and docx:
            try:
                doc = docx.Document(path)
                return "\n".join([p.text for p in doc.paragraphs])
            except:
                return "[Could not read DOCX. Please ensure python-docx is installed or type manually.]"
        elif ext == ".pdf" and PyPDF2:
            try:
                reader = PyPDF2.PdfReader(path)
                return "\n".join([page.extract_text() for page in reader.pages])
            except:
                return "[Could not read PDF. Please ensure PyPDF2 is installed or type manually.]"
        else:
            return f"[File '{os.path.basename(path)}' selected. Text extraction not supported for this format. Please type your message below.]\n\n"

    def proceed(self):
        if self.channel_name == "Email":
            if not self.email_entry.get() or not self.pass_entry.get():
                messagebox.showerror("Error", "Please enter Email and Password.", parent=self)
                return
        
        if not self.text_widget.get("1.0", tk.END).strip() and not self.attachment_path:
            messagebox.showerror("Error", "Please provide a message or select a file.", parent=self)
            return
            
        config = {
            "message": self.text_widget.get("1.0", tk.END).strip(),
            "attachment": self.attachment_path,
            "channel": self.channel_name
        }
        
        if self.channel_name == "Email":
            config["provider"] = self.provider_var.get()
            config["sender"] = self.email_entry.get()
            config["password"] = self.pass_entry.get()
            config["subject"] = self.subject_entry.get()
            config["attach_file"] = self.attach_var.get()
            
        self.destroy()
        threading.Thread(target=self.app.send_messages, args=(config,), daemon=True).start()


# ==============================================================================
#  MAIN APPLICATION WINDOW
# ==============================================================================
class MessageSenderApp:
    def __init__(self, root):
        self.root = root
        self.root.title("Universal Message Sender")
        self.root.geometry("850x650")
        self.root.configure(bg="#1E1E1E")
        
        self.df = None
        self.file_path = None
        self.filtered_df = pd.DataFrame()
        
        self.setup_ui()
        
    def setup_ui(self):
        tk.Label(self.root, text="📨 Universal Message Sender", 
                 font=("Arial", 16, "bold"), fg="#FFF", bg="#1E1E1E").pack(pady=10)
                 
        # File & Filters
        top_frame = tk.Frame(self.root, bg="#1E1E1E")
        top_frame.pack(fill="x", padx=20, pady=5)
        
        self.file_label = tk.Label(top_frame, text="❌ No file loaded", fg="#F44336", bg="#1E1E1E", font=("Arial", 10, "bold"))
        self.file_label.pack(side="left")
        
        tk.Button(top_frame, text="📂 Load Excel", command=self.load_file, 
                  bg="#4CAF50", fg="#FFF", font=("Arial", 10, "bold")).pack(side="left", padx=10)
                  
        filter_frame = tk.Frame(self.root, bg="#1E1E1E")
        filter_frame.pack(fill="x", padx=20, pady=5)
        
        tk.Label(filter_frame, text="Filter by Status:", fg="#FFF", bg="#1E1E1E").pack(side="left")
        self.status_var = tk.StringVar()
        self.status_combo = ttk.Combobox(filter_frame, textvariable=self.status_var, state="readonly", width=15)
        self.status_combo.pack(side="left", padx=5)
        self.status_combo.bind("<<ComboboxSelected>>", lambda e: self.apply_filter())
        
        tk.Label(filter_frame, text="Filter by Group:", fg="#FFF", bg="#1E1E1E").pack(side="left", padx=(20,0))
        self.group_var = tk.StringVar()
        self.group_combo = ttk.Combobox(filter_frame, textvariable=self.group_var, state="readonly", width=15)
        self.group_combo.pack(side="left", padx=5)
        self.group_combo.bind("<<ComboboxSelected>>", lambda e: self.apply_filter())
        
        self.count_label = tk.Label(filter_frame, text="Recipients: 0", fg="#FFA", bg="#1E1E1E", font=("Arial", 10, "bold"))
        self.count_label.pack(side="right")
        
        # Main Buttons
        btn_frame = tk.Frame(self.root, bg="#1E1E1E")
        btn_frame.pack(pady=20)
        
        self.sms_btn = tk.Button(btn_frame, text="📱 Send SMS", command=lambda: self.open_popup("SMS"), 
                                 bg="#1565C0", fg="#FFF", font=("Arial", 12, "bold"), width=18, height=2, state="disabled")
        self.sms_btn.grid(row=0, column=0, padx=10)
        
        self.wa_btn = tk.Button(btn_frame, text="💬 Send WhatsApp", command=lambda: self.open_popup("WhatsApp"), 
                                bg="#2E7D32", fg="#FFF", font=("Arial", 12, "bold"), width=18, height=2, state="disabled")
        self.wa_btn.grid(row=0, column=1, padx=10)
        
        self.email_btn = tk.Button(btn_frame, text="📧 Send Email", command=lambda: self.open_popup("Email"), 
                                   bg="#7B1FA2", fg="#FFF", font=("Arial", 12, "bold"), width=18, height=2, state="disabled")
        self.email_btn.grid(row=0, column=2, padx=10)
        
        tk.Button(btn_frame, text="🚪 Exit", command=self.root.destroy, 
                  bg="#C62828", fg="#FFF", font=("Arial", 12, "bold"), width=18, height=2).grid(row=1, column=0, columnspan=3, pady=15)
                  
        # Log Area
        self.log_text = scrolledtext.ScrolledText(self.root, height=10, bg="#0D1117", fg="#C9D1D9", font=("Consolas", 9), state="disabled")
        self.log_text.pack(fill="both", expand=True, padx=20, pady=10)
        
    def log(self, msg):
        def _update():
            self.log_text.config(state="normal")
            self.log_text.insert(tk.END, msg + "\n")
            self.log_text.see(tk.END)
            self.log_text.config(state="disabled")
        self.root.after(0, _update)
        
    def load_file(self):
        path = filedialog.askopenfilename(filetypes=[("Excel Files", "*.xlsx *.xls")])
        if not path: return
        
        try:
            self.df = pd.read_excel(path)
            self.file_path = path
            
            # --- SMART COLUMN MAPPING ---
            col_map = {}
            for c in self.df.columns:
                c_lower = str(c).lower().strip()
                if c_lower in ['phone', 'telefono', 'tel', 'celular', 'mobile', 'phone number', 'edited phone']:
                    col_map[c] = 'phone'
                elif c_lower in ['email', 'correo', 'e-mail', 'mail', 'correo electronico']:
                    col_map[c] = 'email'
                elif c_lower in ['name', 'nombre', 'nombre completo']:
                    col_map[c] = 'name'
                elif c_lower in ['num', 'numero', 'id', 'no', 'no.', '#']:
                    col_map[c] = 'num'
                elif c_lower in ['group', 'grupo']:
                    col_map[c] = 'group'
                elif c_lower in ['status', 'estado']:
                    col_map[c] = 'status'
                elif c_lower in ['active', 'activo']:
                    col_map[c] = 'active'
                    
            self.df.rename(columns=col_map, inplace=True)
            
            # Ensure required columns exist to prevent crashes
            for col in ['num', 'name', 'phone', 'email', 'status', 'active', 'group']:
                if col not in self.df.columns:
                    self.df[col] = ""
            
            # Add datetime tracking columns if they don't exist
            for col in ['last_sent_sms', 'last_sent_wa', 'last_sent_email']:
                if col not in self.df.columns:
                    self.df[col] = ""
                # CRITICAL FIX: Force column to 'object' (text) type. 
                # If left as 'float64', Pandas will crash when trying to save a text date.
                self.df[col] = self.df[col].astype(object)
            
            self.file_label.config(text=f"✅ {os.path.basename(path)}", fg="#66BB6A")
            
            # Populate filters
            self.status_combo['values'] = ['All'] + sorted(self.df['status'].dropna().unique().tolist())
            self.status_var.set('All')
            
            self.group_combo['values'] = ['All'] + sorted(self.df['group'].dropna().unique().tolist())
            self.group_var.set('All')
            
            self.apply_filter()
            
            self.sms_btn.config(state="normal")
            self.wa_btn.config(state="normal")
            self.email_btn.config(state="normal")
            
            self.log(f"Loaded {len(self.df)} records. Columns mapped successfully.")
            
        except Exception as e:
            messagebox.showerror("Error", f"Failed to load file:\n{e}")
            
    def apply_filter(self):
        if self.df is None: return
        
        status = self.status_var.get()
        group = self.group_var.get()
        
        filtered = self.df.copy()
        if status != 'All':
            filtered = filtered[filtered['status'].astype(str).str.lower() == status.lower()]
        if group != 'All':
            filtered = filtered[filtered['group'].astype(str).str.lower() == group.lower()]
            
        self.filtered_df = filtered
        self.count_label.config(text=f"Recipients: {len(filtered)}")
        
    def get_filtered_df(self):
        return self.filtered_df
        
    def open_popup(self, channel):
        if self.df is None:
            messagebox.showwarning("Warning", "Please load an Excel file first.")
            return
        if len(self.get_filtered_df()) == 0:
            messagebox.showwarning("Warning", "No recipients match the current filters.")
            return
            
        ChannelPopup(self.root, self, channel)
        
    def send_messages(self, config):
        channel = config['channel']
        df = self.filtered_df.copy()
        message = config['message']
        attachment = config.get('attachment')
        
        self.log(f"--- Starting {channel} transmission for {len(df)} recipients ---")
        
        success_count = 0
        col_name = f"last_sent_{'sms' if channel == 'SMS' else 'wa' if channel == 'WhatsApp' else 'email'}"
        
        server = None
        if channel == "Email":
            provider = config['provider']
            smtp_server = "smtp.gmail.com" if provider == "Gmail" else "smtp-mail.outlook.com"
            try:
                server = smtplib.SMTP(smtp_server, 587)
                server.starttls()
                server.login(config['sender'], config['password'])
                self.log("✅ SMTP Login successful")
            except Exception as e:
                self.log(f"❌ SMTP Login failed: {e}")
                return
                
        for index, row in df.iterrows():
            name = str(row.get('name', ''))
            
            # FIX: Handle phone numbers stored as floats in Excel (e.g. 3172742585.0)
            raw_phone = str(row.get('phone', '')).strip()
            if raw_phone.endswith('.0'):
                raw_phone = raw_phone[:-2]
            clean_phone = re.sub(r'\D', '', raw_phone)
            
            email = str(row.get('email', '')).strip()
            
            # Personalize the message for this specific row
            final_message = format_message(message, row)
            
            try:
                if channel == "SMS":
                    if not clean_phone: 
                        self.log(f"⚠️ Skipped {name}: No valid phone number.")
                        continue
                    url = f"sms:{clean_phone}?&body={urllib.parse.quote(final_message)}"
                    webbrowser.open(url)
                    time.sleep(6) # Increased delay so browser can process tabs
                    
                elif channel == "WhatsApp":
                    if not clean_phone: 
                        self.log(f"⚠️ Skipped {name}: No valid phone number.")
                        continue
                    url = f"https://wa.me/{clean_phone}?text={urllib.parse.quote(final_message)}"
                    webbrowser.open(url)
                    time.sleep(10) # Increased delay for WhatsApp Web to load
                    
                elif channel == "Email":
                    if not email or "@" not in email: 
                        self.log(f"⚠️ Skipped {name}: No valid email address.")
                        continue
                    
                    msg = MIMEMultipart()
                    msg['From'] = config['sender']
                    msg['To'] = email
                    final_subject = format_message(config['subject'], row)
                    msg['Subject'] = final_subject
                    msg.attach(MIMEText(final_message, 'plain', 'utf-8'))
                    
                    if config.get('attach_file') and attachment and os.path.exists(attachment):
                        with open(attachment, "rb") as f:
                            part = MIMEBase('application', 'octet-stream')
                            part.set_payload(f.read())
                        encoders.encode_base64(part)
                        part.add_header('Content-Disposition', f"attachment; filename={os.path.basename(attachment)}")
                        msg.attach(part)
                        
                    server.sendmail(config['sender'], email, msg.as_string())
                    time.sleep(8)
                    
                # Update DataFrame with Timestamp (Now safe because column is 'object' type)
                self.df.at[index, col_name] = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
                success_count += 1
                target_info = email if channel == 'Email' else clean_phone
                self.log(f"✅ [{channel}] Sent to {name} ({target_info})")
                
            except Exception as e:
                self.log(f"❌ [{channel}] Failed for {name}: {e}")
                
        if channel == "Email" and server:
            server.quit()
            
        # Save Excel file with new timestamps
        try:
            self.df.to_excel(self.file_path, index=False)
            self.log(f"--- Transmission Complete! {success_count} sent. Excel file updated with timestamps. ---")
            if channel in ["SMS", "WhatsApp"]:
                self.log(f"💡 NOTE: Check your browser's tab bar at the top. The browser hides older tabs behind the last one opened.")
        except Exception as e:
            self.log(f"❌ Failed to save Excel file: {e}. (Make sure the file is not open in Excel).")


# ==============================================================================
#  🚀 LAUNCH TRIGGER
# ==============================================================================
if __name__ == "__main__":
    root = tk.Tk()
    app = MessageSenderApp(root)
    root.mainloop()